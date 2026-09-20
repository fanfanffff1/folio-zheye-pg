#!/usr/bin/env python3
"""Import 华语文学藏书 (Chinese) from ../华语文学藏书 into the Book table.

- reads 2026华语文学藏书清单.xlsx (same columns as the other catalogs)
- copies covers from 封面/ into static/covers/
- creates/updates books with language_code="zh"

Idempotent: matched by slug.

Usage:
    python3 scripts/import_huayu.py
    DATABASE_URL="<neon>" python3 scripts/import_huayu.py   # import to production
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INSPIRATION = ROOT.parent
FOLDER = INSPIRATION / "华语文学藏书"
XLSX = FOLDER / "2026华语文学藏书清单.xlsx"
COVERS_SRC = FOLDER / "封面"
COVERS_DST = ROOT / "static" / "covers"

sys.path.insert(0, str(ROOT))

from folio.cover_urls import derive_cover_urls  # noqa: E402
from folio.models import Author, Book, SessionLocal, init_db  # noqa: E402


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"找不到表格：{XLSX}")
    init_db()
    wb = load_workbook(XLSX, read_only=True)
    ws = wb["全部书目"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    idx = {h: i for i, h in enumerate(hdr)}

    db = SessionLocal()
    n = 0
    try:
        for r in it:
            def g(key: str) -> str:
                i = idx.get(key)
                return str(r[i]).strip() if i is not None and r[i] is not None else ""

            slug = g("slug")
            if not slug:
                continue
            title = g("中文常用译名") or g("原文书名")
            author_name = g("作者")

            author = None
            if author_name:
                author = db.query(Author).filter(Author.name == author_name).one_or_none()
                if not author:
                    author = Author(name=author_name, localized_name=author_name)
                    db.add(author)
                    db.flush()

            cover_image = "/covers/placeholder.svg"
            cover_file = g("封面文件")
            if cover_file and (COVERS_SRC / cover_file).exists():
                COVERS_DST.mkdir(parents=True, exist_ok=True)
                dst = COVERS_DST / (slug + Path(cover_file).suffix)
                shutil.copyfile(COVERS_SRC / cover_file, dst)
                cover_image = "/covers/" + dst.name

            book = db.query(Book).filter(Book.slug == slug).one_or_none()
            if not book:
                book = Book(slug=slug)
                db.add(book)
            book.original_title = g("原文书名") or title
            book.chinese_title = title
            book.author_id = author.id if author else None
            book.language_code = "zh"
            book.language_name = "中文"
            book.publisher = g("原版代表出版社")
            yr = g("首次出版年")[:4]
            book.publication_year = int(yr) if yr.isdigit() else None
            book.isbn13 = g("ISBN-13")
            book.primary_genre = g("主类型") or "文学"
            book.short_description_zh = g("一句话简介")
            book.full_description_zh = g("故事梗概")
            book.recommendation_zh = g("入选理由")
            book.audience_zh = g("适合谁")
            book.cover_image = cover_image
            thumb, full = derive_cover_urls(cover_image)
            book.cover_thumbnail_url = thumb
            book.cover_full_url = full
            book.source_file = XLSX.name
            n += 1
        db.commit()
        print(f"imported {n} books from {XLSX.name}")
    finally:
        db.close()
        wb.close()


if __name__ == "__main__":
    main()
