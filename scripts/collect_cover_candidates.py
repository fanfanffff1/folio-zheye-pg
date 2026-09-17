#!/usr/bin/env python3
"""Step 1: collect HD cover *candidates* for FR/KO books that still use SVG title cards.

Saves into the language folders (outside deploy path) for manual screening:
  ../法语原版书籍/封面搜集/候选/{slug}.jpg
  ../韩语原版书籍/封面搜集/封面整理进度.xlsx

Does NOT modify folio DB or R2. Later steps: screen → optimize → upload → wire paths.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from folio.models import Book, SessionLocal  # noqa: E402

INSPIRATION = ROOT.parent
LANG_FOLDERS = {
    "fr": INSPIRATION / "法语原版书籍" / "封面搜集",
    "ko": INSPIRATION / "韩语原版书籍" / "封面搜集",
}
UA = "FOLIO-cover-collect/1.0 (editorial; local research)"


def http_get(url: str, binary: bool = False, timeout: int = 25):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    last_err = None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                data = resp.read()
            break
        except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, ConnectionError) as exc:
            last_err = exc
            time.sleep(0.6 * (attempt + 1))
            data = None
        except Exception as exc:  # RemoteDisconnected etc.
            last_err = exc
            time.sleep(0.6 * (attempt + 1))
            data = None
    else:
        return None
    if data is None:
        return None
    if binary:
        return data
    try:
        return json.loads(data.decode("utf-8", "replace"))
    except json.JSONDecodeError:
        return None


def save_cover_bytes(data: bytes, dest: Path, min_side: int = 200) -> bool:
    try:
        im = Image.open(BytesIO(data))
        im = ImageOps.exif_transpose(im)
    except Exception:
        return False
    if min(im.size) < min_side:
        return False
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        bg = Image.new("RGB", im.size, (255, 255, 255))
        rgba = im.convert("RGBA")
        bg.paste(rgba, mask=rgba.split()[-1])
        im = bg
    else:
        im = im.convert("RGB")
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, format="JPEG", quality=92, optimize=True)
    return dest.is_file() and dest.stat().st_size > 8_000


def openlibrary_by_isbn(isbn: str) -> bytes | None:
    isbn = re.sub(r"\D", "", isbn or "")
    if len(isbn) not in (10, 13):
        return None
    for size in ("L", "M"):
        data = http_get(f"https://covers.openlibrary.org/b/isbn/{isbn}-{size}.jpg", binary=True)
        if data and len(data) > 8_000 and data[:3] != b"GIF":
            return data
    return None


def openlibrary_search(title: str, author: str, lang: str) -> tuple[bytes | None, str]:
    q = title
    if author:
        q = f"{title} {author}"
    params = urllib.parse.urlencode({"title": title, "author": author or "", "limit": 5})
    payload = http_get(f"https://openlibrary.org/search.json?{params}")
    if not payload:
        return None, ""
    for doc in payload.get("docs") or []:
        langs = doc.get("language") or []
        want = {"fr": "fre", "ko": "kor"}.get(lang)
        if want and langs and want not in langs:
            # still allow if title is close enough / no lang
            pass
        cover_id = doc.get("cover_i")
        isbn_list = doc.get("isbn") or []
        data = None
        source = ""
        if isbn_list:
            data = openlibrary_by_isbn(isbn_list[0])
            source = f"OpenLibrary ISBN {isbn_list[0]}"
        if not data and cover_id:
            data = http_get(f"https://covers.openlibrary.org/b/id/{cover_id}-L.jpg", binary=True)
            source = f"OpenLibrary cover_i={cover_id}"
        if data and len(data) > 8_000:
            return data, source
    return None, ""


def google_books_cover(title: str, author: str, lang: str, isbn: str = "") -> tuple[bytes | None, str]:
    if isbn:
        isbn_digits = re.sub(r"\D", "", isbn)
        q = f"isbn:{isbn_digits}"
    else:
        q = f'intitle:"{title}"'
        if author:
            token = author.split()[0]
            q += f" inauthor:{token}"
    params = {"q": q, "maxResults": "5"}
    if lang in {"fr", "ko", "en", "es", "it", "ja"}:
        params["langRestrict"] = lang
    url = "https://www.googleapis.com/books/v1/volumes?" + urllib.parse.urlencode(params)
    payload = http_get(url)
    if not payload:
        return None, ""
    for item in payload.get("items") or []:
        info = item.get("volumeInfo") or {}
        image_links = info.get("imageLinks") or {}
        for key in ("extraLarge", "large", "medium", "thumbnail", "smallThumbnail"):
            raw = image_links.get(key)
            if not raw:
                continue
            raw = raw.replace("http://", "https://").replace("&edge=curl", "")
            if "zoom=" in raw:
                raw = re.sub(r"zoom=\d", "zoom=3", raw)
            data = http_get(raw, binary=True)
            if data and len(data) > 5_000:
                return data, f"Google Books ({key})"
        for ident in info.get("industryIdentifiers") or []:
            if ident.get("type") in {"ISBN_13", "ISBN_10"}:
                num = re.sub(r"\D", "", ident.get("identifier") or "")
                data = http_get(
                    f"https://books.google.com/books?vid=ISBN{num}&printsec=frontcover&img=1&zoom=4",
                    binary=True,
                )
                if data and len(data) > 5_000:
                    return data, f"Google Books ISBN {num}"
    return None, ""


def needs_photo(book: Book) -> bool:
    c = (book.cover_image or "").strip()
    if not c or "placeholder" in c:
        return True
    return "/covers/card-" in c or c.endswith(".svg")


def ensure_workbook(path: Path) -> None:
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "封面整理"
    ws.append(
        [
            "序号",
            "slug",
            "原文书名",
            "中文名",
            "作者",
            "出版年",
            "ISBN",
            "当前封面",
            "候选文件",
            "候选来源",
            "筛选结果",
            "最终文件名",
            "备注",
        ]
    )
    wb.save(path)


def upsert_row(ws, slug: str, values: dict) -> None:
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    found = None
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, col["slug"]).value == slug:
            found = row
            break
    if found is None:
        found = ws.max_row + 1
        ws.cell(found, col["序号"], found - 1)
        ws.cell(found, col["slug"], slug)
    for key, val in values.items():
        if key in col:
            ws.cell(found, col[key], val)


def collect_lang(lang: str, limit: int = 0, sleep_s: float = 0.35) -> None:
    base = LANG_FOLDERS[lang]
    cand_dir = base / "候选"
    done_dir = base / "已通过"
    cand_dir.mkdir(parents=True, exist_ok=True)
    done_dir.mkdir(parents=True, exist_ok=True)
    xlsx = base / "封面整理进度.xlsx"
    ensure_workbook(xlsx)
    wb = load_workbook(xlsx)
    ws = wb["封面整理"]

    db = SessionLocal()
    books = (
        db.query(Book)
        .filter(Book.language_code == lang)
        .order_by(Book.id)
        .all()
    )
    todo = [b for b in books if needs_photo(b)]
    if limit:
        todo = todo[:limit]
    print(f"[{lang}] need photo={len(todo)} / total={len(books)} → {cand_dir}")

    ok = fail = skip = 0
    for i, book in enumerate(todo, 1):
        dest = cand_dir / f"{book.slug}.jpg"
        author = book.author.name if book.author else ""
        year = book.publication_year or ""
        isbn = (book.isbn13 or book.isbn10 or "").strip()
        current = book.cover_image or ""
        if dest.exists() and dest.stat().st_size > 8_000:
            skip += 1
            upsert_row(
                ws,
                book.slug,
                {
                    "原文书名": book.original_title,
                    "中文名": book.chinese_title,
                    "作者": author,
                    "出版年": year,
                    "ISBN": isbn,
                    "当前封面": current,
                    "候选文件": str(dest.relative_to(base)),
                    "候选来源": "已存在本地候选",
                    "筛选结果": "待筛",
                },
            )
            print(f"  [{i}/{len(todo)}] skip existing {book.slug}")
            continue

        data = None
        source = ""
        if isbn:
            data = openlibrary_by_isbn(isbn)
            if data:
                source = f"OpenLibrary ISBN {isbn}"
        if not data:
            data, source = openlibrary_search(book.original_title or "", author, lang)
        if not data:
            data, source = google_books_cover(book.original_title or "", author, lang, isbn)
            time.sleep(sleep_s)

        if data and save_cover_bytes(data, dest):
            ok += 1
            status = "待筛"
            note = ""
            print(f"  [{i}/{len(todo)}] OK {book.slug} ← {source} ({dest.stat().st_size // 1024}KB)")
        else:
            fail += 1
            status = "未找到"
            note = "自动检索未命中，需手工搜图"
            source = source or ""
            if dest.exists():
                try:
                    dest.unlink()
                except OSError:
                    pass
            print(f"  [{i}/{len(todo)}] FAIL {book.slug}")

        upsert_row(
            ws,
            book.slug,
            {
                "原文书名": book.original_title,
                "中文名": book.chinese_title,
                "作者": author,
                "出版年": year,
                "ISBN": isbn,
                "当前封面": current,
                "候选文件": str(dest.relative_to(base)) if status == "待筛" else "",
                "候选来源": source,
                "筛选结果": status,
                "备注": note,
            },
        )
        if i % 10 == 0:
            wb.save(xlsx)
        time.sleep(sleep_s)

    wb.save(xlsx)
    db.close()
    print(f"[{lang}] done ok={ok} fail={fail} skip={skip} sheet={xlsx}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--lang", choices=("fr", "ko", "all"), default="all")
    parser.add_argument("--limit", type=int, default=0, help="max books per language (0=all)")
    args = parser.parse_args()
    langs = ["fr", "ko"] if args.lang == "all" else [args.lang]
    for lang in langs:
        collect_lang(lang, limit=args.limit)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
