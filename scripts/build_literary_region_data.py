#!/usr/bin/env python3
"""Build folio/data/literary_region_books.json from the source xlsx.

Source: 中国地域与世界华文文学百年书单.xlsx
  - 地区脉络   : 百年发展脉络 / 核心风格 / 范围与交叉说明
  - 作者作品库 : 地区 / 时期层 / 作者 / 代表作1-3 / 重要作4-5 / 纳入理由

Usage:
  python3 scripts/build_literary_region_data.py [--xlsx PATH]
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path("/Users/zhangfanfan/Desktop/inspiration/中国地域与世界华文文学百年书单.xlsx")
OUT = ROOT / "folio" / "data" / "literary_region_books.json"

# xlsx 地区名 → 地图 region id（与 region_groups.py 对齐）
NAME_TO_ID = {
    "西北": "xibei",
    "藏地": "zangdi",
    "西南": "xinan",
    "中原": "zhongyuan",
    "华北京津": "huabei",
    "江南": "jiangnan",
    "东南·闽台": "dongnan",
    "岭南": "lingnan",
    "草原": "caoyuan",
    "东北": "dongbei",
    "香港": "hongkong",
    "澳门": "macau",
    "台湾": "taiwan",
    "北美华语": "northamerica",
    "欧洲华语": "europe",
    "日韩华语": "japankorea",
    "东南亚华语": "seasia",
    "马华": "mahua",
    "新加坡华文": "singapore",
    "大洋洲华语": "oceania",
}

PERIOD_KEY = {"历史/先驱": "historical", "当代": "contemporary"}
PHASE_SPLIT = re.compile(r"[；;]")


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_phases(context: str) -> list[dict[str, str]]:
    phases = []
    for i, chunk in enumerate(PHASE_SPLIT.split(context or "")):
        chunk = chunk.strip()
        if not chunk:
            continue
        if "：" in chunk:
            years, _, text = chunk.partition("：")
        elif ":" in chunk:
            years, _, text = chunk.partition(":")
        else:
            years, text = "", chunk
        phases.append({"id": f"p{i + 1}", "years": years.strip(), "text": text.strip()})
    return phases


def read_context(wb) -> dict[str, dict]:
    ws = wb["地区脉络"]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        name = clean(row[0])
        if name not in NAME_TO_ID:
            continue
        context = clean(row[1])
        out[NAME_TO_ID[name]] = {
            "region_name": name,
            "context": context,
            "phases": parse_phases(context),
            "style": clean(row[2]),
            "scope": clean(row[3]),
        }
    return out


def read_authors(wb) -> dict[str, dict[str, list[dict]]]:
    ws = wb["作者作品库"]
    out: dict[str, dict[str, list[dict]]] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        region_name = clean(row[1])
        if region_name not in NAME_TO_ID:
            continue
        rid = NAME_TO_ID[region_name]
        period = PERIOD_KEY.get(clean(row[2]), "contemporary")
        name = clean(row[3])
        if not name:
            continue
        works: list[str] = []
        for cell in row[4:9]:
            title = clean(cell)
            if title and title not in works:
                works.append(title)
        entry = {
            "name": name,
            "works": works,
            "reason": clean(row[9]),
            "cross": clean(row[10]),
        }
        out.setdefault(rid, {"historical": [], "contemporary": []})[period].append(entry)
    return out


def pick_representatives(authors: dict[str, list[dict]], limit: int = 4) -> list[str]:
    hist = authors.get("historical", [])
    cont = authors.get("contemporary", [])
    half = max(1, limit // 2)
    picks = [a["name"] for a in hist[:half]] + [a["name"] for a in cont[:half]]
    seen, out = set(), []
    for name in picks:
        if name not in seen:
            seen.add(name)
            out.append(name)
    return out


def pick_works(authors: dict[str, list[dict]], limit: int = 6) -> list[dict[str, str]]:
    hist = authors.get("historical", [])
    cont = authors.get("contemporary", [])
    half = max(1, limit // 2)
    ordered = hist[:half] + cont[:half]
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for author in ordered:
        for title in author["works"][:1]:
            if title in seen:
                continue
            seen.add(title)
            out.append({"title": title, "author": author["name"]})
    # 若不足，补足后面的作者
    for author in hist[half:] + cont[half:]:
        if len(out) >= limit:
            break
        for title in author["works"][:1]:
            if title in seen:
                continue
            seen.add(title)
            out.append({"title": title, "author": author["name"]})
    return out[:limit]


def build(xlsx: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    contexts = read_context(wb)
    author_map = read_authors(wb)

    regions: dict[str, dict] = {}
    for rid, ctx in contexts.items():
        authors = author_map.get(rid, {"historical": [], "contemporary": []})
        regions[rid] = {
            **ctx,
            "authors": authors,
            "representative_authors": pick_representatives(authors),
            "representative_works": pick_works(authors),
            "image": f"/static/img/literary-map/region-{rid}.jpg",
        }
    return {"source": xlsx.name, "regions": regions}


def main() -> int:
    parser = argparse.ArgumentParser(description="Build literary region data JSON")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    args = parser.parse_args()
    if not args.xlsx.exists():
        raise SystemExit(f"xlsx not found: {args.xlsx}")
    data = build(args.xlsx)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    total_authors = sum(
        len(r["authors"]["historical"]) + len(r["authors"]["contemporary"])
        for r in data["regions"].values()
    )
    print(f"wrote {OUT} regions={len(data['regions'])} authors={total_authors}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
